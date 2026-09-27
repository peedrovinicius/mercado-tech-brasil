from __future__ import annotations

import hashlib
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from src.api.publication import published_competencies


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def normalize_cbo_code(value: object) -> str | None:
    if value is None:
        return None

    if isinstance(value, bool):
        return None

    if isinstance(value, int):
        text = str(value)
    elif isinstance(value, float):
        if not value.is_integer():
            return None
        text = str(int(value))
    else:
        text = "".join(character for character in str(value) if character.isdigit())

    return text if len(text) == 6 else None


def published_occupation_codes(gold_dir: Path) -> set[str]:
    codes: set[str] = set()

    for competence in published_competencies(gold_dir):
        path = gold_dir / f"by-occupation-{competence}.json"
        if not path.exists():
            raise FileNotFoundError(
                f"Artefato ocupacional publicado ausente: {path}"
            )

        payload = json.loads(path.read_text(encoding="utf-8"))
        raw_items = payload.get("items")
        if not isinstance(raw_items, list):
            raise TypeError(f"Lista ocupacional inválida em {path.name}.")

        for item in raw_items:
            if not isinstance(item, dict):
                raise TypeError(f"Item ocupacional inválido em {path.name}.")
            code = normalize_cbo_code(item.get("cbo_codigo"))
            if code is None:
                raise ValueError(
                    f"Código CBO inválido no Gold publicado: {item.get('cbo_codigo')!r}."
                )
            codes.add(code)

    if not codes:
        raise ValueError("Nenhum código CBO tech publicado foi encontrado.")

    return codes


def inspect_qbq_workbook(
    path: Path,
    *,
    expected_codes: set[str],
    max_scan_rows: int = 10000,
) -> dict[str, Any]:
    if path.suffix.lower() != ".xlsx":
        raise ValueError("O arquivo QBQ precisa estar no formato .xlsx.")
    if not path.exists():
        raise FileNotFoundError(path)
    if not expected_codes:
        raise ValueError("Nenhum código CBO esperado foi informado.")
    if max_scan_rows < 1:
        raise ValueError("max_scan_rows precisa ser maior que zero.")

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheets: list[dict[str, Any]] = []
    all_matched_codes: set[str] = set()
    detected_join_columns: list[dict[str, Any]] = []

    try:
        for worksheet in workbook.worksheets:
            max_column = min(int(worksheet.max_column or 0), 200)
            max_row = min(int(worksheet.max_row or 0), max_scan_rows)
            headers: dict[int, str] = {}
            matches_by_column: dict[int, set[str]] = {}

            for row_index, row in enumerate(
                worksheet.iter_rows(
                    min_row=1,
                    max_row=max_row,
                    max_col=max_column,
                    values_only=True,
                ),
                start=1,
            ):
                if row_index == 1:
                    headers = {
                        index: str(value).strip()
                        for index, value in enumerate(row, start=1)
                        if value is not None and str(value).strip()
                    }

                for column_index, value in enumerate(row, start=1):
                    code = normalize_cbo_code(value)
                    if code is None or code not in expected_codes:
                        continue
                    matches_by_column.setdefault(column_index, set()).add(code)
                    all_matched_codes.add(code)

            join_columns: list[dict[str, Any]] = []
            for column_index, matched in sorted(matches_by_column.items()):
                item = {
                    "column_index": column_index,
                    "header": headers.get(column_index),
                    "matched_codes": sorted(matched),
                    "matched_count": len(matched),
                }
                join_columns.append(item)
                detected_join_columns.append(
                    {
                        "sheet": worksheet.title,
                        **item,
                    }
                )

            sheets.append(
                {
                    "name": worksheet.title,
                    "rows_reported": int(worksheet.max_row or 0),
                    "columns_reported": int(worksheet.max_column or 0),
                    "scanned_rows": max_row,
                    "headers": [
                        {
                            "column_index": index,
                            "value": value,
                        }
                        for index, value in sorted(headers.items())
                    ],
                    "detected_join_columns": join_columns,
                }
            )
    finally:
        workbook.close()

    missing = sorted(expected_codes - all_matched_codes)
    return {
        "version": 1,
        "source_file": path.name,
        "source_sha256": _sha256(path),
        "source_size_bytes": path.stat().st_size,
        "expected_tech_codes": sorted(expected_codes),
        "matched_tech_codes": sorted(all_matched_codes),
        "missing_tech_codes": missing,
        "coverage_complete": not missing,
        "detected_join_columns": detected_join_columns,
        "sheets": sheets,
    }


def validate_qbq_workbook(
    path: Path,
    *,
    gold_dir: Path,
    max_scan_rows: int = 10000,
) -> dict[str, Any]:
    expected_codes = published_occupation_codes(gold_dir)
    report = inspect_qbq_workbook(
        path,
        expected_codes=expected_codes,
        max_scan_rows=max_scan_rows,
    )

    detected = report["detected_join_columns"]
    if not isinstance(detected, list) or not detected:
        raise ValueError(
            "Nenhuma coluna do workbook pôde ser ligada aos códigos CBO tech publicados."
        )

    return report


def write_qbq_inspection_report(report: dict[str, Any], destination: Path) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(
        json.dumps(report, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
